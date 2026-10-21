import os
import sys
import io
import time
import unittest

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from app import app
from config import Config
from docx import Document

class TestAPIIntegration(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_get_index_page(self):
        res = self.client.get('/')
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"DocuShift AI", res.data)
        self.assertIn(b"Editable Mode", res.data)
        self.assertIn(b"High-Fidelity Mode", res.data)

    def test_convert_with_mode_selection(self):
        sample_path = os.path.join(BASE_DIR, 'tests', 'fixtures', 'sample_resume.jpg')
        self.assertTrue(os.path.exists(sample_path), "Sample image fixture must exist for testing")

        with open(sample_path, 'rb') as f:
            img_bytes = f.read()

        # Test POST /convert with mode='high_fidelity'
        data = {
            'document': (io.BytesIO(img_bytes), 'test_upload.jpg'),
            'mode': 'high_fidelity'
        }
        res = self.client.post('/convert', data=data)
        self.assertEqual(res.status_code, 202)
        json_data = res.get_json()
        self.assertTrue(json_data['success'])
        self.assertEqual(json_data['mode'], 'high_fidelity')
        job_id = json_data['job_id']

        # Poll status (allow up to 70s for network & Gemini API backoff retries)
        max_retries = 35
        completed = False
        for _ in range(max_retries):
            time.sleep(2)
            st_res = self.client.get(f'/status/{job_id}')
            self.assertEqual(st_res.status_code, 200)
            st_json = st_res.get_json()
            if st_json['status'] == 'completed':
                completed = True
                self.assertIsNotNone(st_json['download_url'])
                break
            elif st_json['status'] == 'failed':
                self.fail(f"Conversion job failed: {st_json.get('error')}")

        self.assertTrue(completed, "Conversion job should complete within timeout")

        # Test GET /download/<job_id>
        dl_res = self.client.get(f'/download/{job_id}')
        self.assertEqual(dl_res.status_code, 200)
        self.assertEqual(
            dl_res.mimetype,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        self.assertGreater(len(dl_res.data), 1000)

        # Verify Word document is readable
        doc = Document(io.BytesIO(dl_res.data))
        total_p = len(doc.paragraphs) + sum(len(c.paragraphs) for t in doc.tables for r in t.rows for c in r.cells)
        self.assertGreater(total_p, 0, "Document should contain reconstructed paragraphs")

if __name__ == '__main__':
    unittest.main()
