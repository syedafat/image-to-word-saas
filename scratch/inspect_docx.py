import os, sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

def inspect_docx(path):
    print("=== Inspecting:", path)
    if not os.path.exists(path):
        print("File not found")
        return
    doc = Document(path)
    print("Paragraphs outside tables:", len(doc.paragraphs))
    for p in doc.paragraphs:
        if p.text.strip():
            print("  Top/Bottom Para:", repr(p.text.strip()[:60]))
    print("Tables:", len(doc.tables))
    for ti, t in enumerate(doc.tables):
        print(f"Table {ti}: {len(t.rows)} rows x {len(t.columns)} cols")
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                txt = cell.text.strip().replace("\n", " | ")
                if txt:
                    print(f"  [{ti}] Cell({ri},{ci}): {repr(txt[:100])}")
                for ni, nested_table in enumerate(cell.tables):
                    print(f"    Nested Table {ni} in Cell({ri},{ci}): {len(nested_table.rows)} rows x {len(nested_table.columns)} cols")
                    for nri, nrow in enumerate(nested_table.rows):
                        for nci, ncell in enumerate(nrow.cells):
                            ntxt = ncell.text.strip().replace("\n", " | ")
                            print(f"      Nested Cell({nri},{nci}): {repr(ntxt[:80])}")

if __name__ == "__main__":
    inspect_docx("outputs/test_rs_reconstructed.docx")
    inspect_docx("outputs/test_bb_reconstructed.docx")
