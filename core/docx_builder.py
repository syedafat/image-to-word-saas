import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

try:
    from PIL import Image as PILImage, ImageDraw as PILDraw
    _HAS_PIL = True
except ImportError:
    _HAS_PIL = False


def make_circular_image(src_path: str, border_width: int = 6,
                        border_color=(255, 255, 255)) -> str:
    """
    Creates a circular-cropped version of the source image with a solid ring
    border, saves it alongside the original with a '_circle' suffix, and
    returns the path to the new file.  Falls back to the original path when
    PIL is unavailable or on any error.
    """
    if not _HAS_PIL:
        return src_path
    try:
        img = PILImage.open(src_path).convert("RGBA")
        # Make it square first (centre-crop the longer axis)
        w, h = img.size
        side = min(w, h)
        left = (w - side) // 2
        top = (h - side) // 2
        img = img.crop((left, top, left + side, top + side))

        # Anti-aliased circle mask (draw at 4x then downscale)
        scale = 4
        big = side * scale
        mask = PILImage.new("L", (big, big), 0)
        PILDraw.Draw(mask).ellipse((0, 0, big - 1, big - 1), fill=255)
        mask = mask.resize((side, side), PILImage.LANCZOS)

        # Apply mask
        result = PILImage.new("RGBA", (side, side), (0, 0, 0, 0))
        result.paste(img, (0, 0), mask)

        # Draw border ring
        if border_width > 0:
            ring = PILImage.new("RGBA", (side, side), (0, 0, 0, 0))
            draw = PILDraw.Draw(ring)
            bw = border_width
            draw.ellipse((0, 0, side - 1, side - 1),
                         outline=border_color + (255,), width=bw)
            result = PILImage.alpha_composite(result, ring)

        # Save as PNG next to original
        out_path = src_path.rsplit(".", 1)[0] + "_circle.png"
        result.save(out_path, "PNG")
        return out_path
    except Exception:
        return src_path

def generate_corner_triangle_asset(color_hex: str, corner: str = "top_left", size_px: int = 180, assets_dir: str = "outputs/assets") -> str:
    """
    Generates a clean anti-aliased corner triangle PNG.
    """
    if not _HAS_PIL:
        return ""
    try:
        os.makedirs(assets_dir, exist_ok=True)
        c_hex = clean_hex_color(color_hex, default="E5A914")
        r, g, b = int(c_hex[0:2], 16), int(c_hex[2:4], 16), int(c_hex[4:6], 16)
        
        scale = 4
        s = size_px * scale
        im = PILImage.new("RGBA", (s, s), (0, 0, 0, 0))
        draw = PILDraw.Draw(im)
        
        if corner == "top_left":
            draw.polygon([(0, 0), (s, 0), (0, s)], fill=(r, g, b, 255))
        elif corner == "bottom_right":
            draw.polygon([(s, 0), (s, s), (0, s)], fill=(r, g, b, 255))
        elif corner == "top_right":
            draw.polygon([(0, 0), (s, 0), (s, s)], fill=(r, g, b, 255))
        elif corner == "bottom_left":
            draw.polygon([(0, 0), (0, s), (s, s)], fill=(r, g, b, 255))
            
        im = im.resize((size_px, size_px), PILImage.LANCZOS)
        out_path = os.path.join(assets_dir, f"corner_{corner}_{c_hex}.png")
        im.save(out_path, "PNG")
        return out_path
    except Exception:
        return ""

def render_skill_bars_grid(container, skill_items: list[dict], available_width_inches: float = 4.8, accent_hex: str = "E5A914"):
    """
    Renders skill progress bars in a clean 2-column or 1-column layout using native Word tables.
    """
    if not skill_items:
        return
        
    accent_clean = clean_hex_color(accent_hex, default="E5A914")
    cols = 2 if len(skill_items) >= 2 else 1
    rows = (len(skill_items) + cols - 1) // cols
    
    table = container.add_table(rows=rows, cols=cols)
    remove_table_borders(table)
    
    col_w = available_width_inches / cols
    bar_w_inches = col_w - 0.25
    
    for idx, item in enumerate(skill_items):
        r_idx = idx // cols
        c_idx = idx % cols
        cell = table.cell(r_idx, c_idx)
        cell.width = Inches(col_w)
        set_cell_margins(cell, top=30, bottom=40, left=40, right=40)
        
        name = item.get("skill_name") or item.get("text") or "Skill"
        pct = max(10, min(100, int(item.get("skill_percent") or 80)))
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(name)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.bold = True
        run.font.color.rgb = RGBColor(51, 65, 85)
        
        # Skill level bar: 1-row, 2-col mini table inside the cell
        bar_tbl = cell.add_table(rows=1, cols=2)
        remove_table_borders(bar_tbl)
        
        trPr = bar_tbl.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:trHeight %s w:val="90" w:hRule="exact"/>' % nsdecls('w')))
        
        fill_w = max(0.1, bar_w_inches * (pct / 100.0))
        track_w = max(0.1, bar_w_inches * ((100 - pct) / 100.0))
        
        c_fill = bar_tbl.cell(0, 0)
        c_track = bar_tbl.cell(0, 1)
        c_fill.width = Inches(fill_w)
        c_track.width = Inches(track_w)
        
        set_cell_margins(c_fill, top=0, bottom=0, left=0, right=0)
        set_cell_margins(c_track, top=0, bottom=0, left=0, right=0)
        
        shd_fill = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), accent_clean))
        c_fill._tc.get_or_add_tcPr().append(shd_fill)
        
        shd_track = parse_xml(r'<w:shd %s w:fill="E2E8F0"/>' % nsdecls('w'))
        c_track._tc.get_or_add_tcPr().append(shd_track)
        
        for c in (c_fill, c_track):
            cp = c.paragraphs[0]
            cp.paragraph_format.line_spacing = 1.0
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            crun = cp.add_run()
            crun.font.size = Pt(1)

def hex_to_rgb(hex_str: str, default=(15, 23, 42)) -> RGBColor:
    """
    Parses a hex color string into a docx RGBColor instance.
    """
    if not hex_str or not isinstance(hex_str, str):
        return RGBColor(*default)
    hex_clean = hex_str.strip().lstrip('#')
    if len(hex_clean) == 3:
        hex_clean = ''.join([c * 2 for c in hex_clean])
    if len(hex_clean) == 6:
        try:
            r = int(hex_clean[0:2], 16)
            g = int(hex_clean[2:4], 16)
            b = int(hex_clean[4:6], 16)
            return RGBColor(r, g, b)
        except ValueError:
            pass
    return RGBColor(*default)

def clean_hex_color(hex_str: str, default="F1F5F9") -> str:
    """
    Cleans and standardizes hex color for WordprocessingML XML (without '#').
    """
    if not hex_str or not isinstance(hex_str, str):
        return default
    clean = hex_str.strip().lstrip('#').upper()
    if len(clean) == 3:
        clean = ''.join([c * 2 for c in clean])
    if len(clean) == 6:
        return clean
    return default

def get_alignment(align_str: str) -> WD_ALIGN_PARAGRAPH:
    """
    Maps text alignment string to docx WD_ALIGN_PARAGRAPH enum.
    """
    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY
    }
    return align_map.get(str(align_str).lower().strip(), WD_ALIGN_PARAGRAPH.LEFT)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """
    Sets inner cell padding (margins) in dxa (twentieths of a point).
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        r'<w:tcMar %s>'
        r'  <w:top w:w="%d" w:type="dxa"/>'
        r'  <w:bottom w:w="%d" w:type="dxa"/>'
        r'  <w:left w:w="%d" w:type="dxa"/>'
        r'  <w:right w:w="%d" w:type="dxa"/>'
        r'</w:tcMar>' % (nsdecls('w'), top, bottom, left, right)
    )
    tcPr.append(tcMar)

def set_table_borders(table, border_color="CBD5E1", inside_h_color="E2E8F0"):
    """
    Applies clean modern borders to a native Word table.
    """
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="6" w:space="0" w:color="%s"/>'
        r'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="%s"/>'
        r'  <w:left w:val="none"/>'
        r'  <w:right w:val="none"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
        r'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
        r'</w:tblBorders>' % (nsdecls('w'), border_color, border_color, inside_h_color, inside_h_color)
    )
    tblPr.append(tblBorders)

def remove_table_borders(table):
    """
    Removes all borders from a layout grid table and ensures no table styles
    or non-printing gridlines add visible borders.
    """
    try:
        table.style = 'Normal Table'
    except Exception:
        pass
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)
    # Also explicitly clear cell borders on every cell
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                r'<w:tcBorders %s>'
                r'  <w:top w:val="none"/>'
                r'  <w:bottom w:val="none"/>'
                r'  <w:left w:val="none"/>'
                r'  <w:right w:val="none"/>'
                r'</w:tcBorders>' % nsdecls('w')
            )
            tcPr.append(tcBorders)

def render_divider(container, block, available_width_inches=7.5):
    """
    Renders an intentional decorative horizontal rule or dividing line.
    Uses compact vertical spacing to avoid pushing content onto subsequent pages.
    """
    color = clean_hex_color(block.get("text_color"), default="94A3B8")
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        r'<w:pBdr %s>'
        r'  <w:bottom w:val="single" w:sz="8" w:space="2" w:color="%s"/>'
        r'</w:pBdr>' % (nsdecls('w'), color)
    )
    pPr.append(pBdr)

def filter_redundant_dividers(elements: list[dict]) -> list[dict]:
    """
    Removes redundant divider elements that immediately follow headings
    (which already have integrated bottom underline borders), preventing
    double divider lines and saving vertical layout space.
    """
    cleaned = []
    for el in elements:
        if el.get("type") == "divider":
            if cleaned and (cleaned[-1].get("type") == "heading" or cleaned[-1].get("is_heading")):
                continue
        cleaned.append(el)
    return cleaned

def render_image_block(container, block, available_width_inches=7.5, mode="editable", max_image_width=2.5):
    """
    Renders an extracted high-resolution image asset (photo, logo, signature).
    Places each image in its correct relative section rather than stacking at top.
    """
    asset_path = block.get("image_asset_path")
    if not asset_path or not os.path.exists(asset_path):
        return

    bbox = block.get("bbox", [])
    if len(bbox) == 4:
        ymin, xmin, ymax, xmax = bbox
        rel_width = max(0.15, min(1.0, (xmax - xmin) / 1000.0))
        target_width_inches = max(0.8, min(max_image_width, available_width_inches * rel_width))
    else:
        target_width_inches = min(available_width_inches, max_image_width)

    p = container.add_paragraph()
    p.alignment = get_alignment(block.get("text_align", "center"))
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)

    # Apply circular mask for profile/portrait photos
    label = str(block.get("image_label", "")).lower()
    if "photo" in label or "profile" in label:
        asset_path = make_circular_image(asset_path, border_width=6, border_color=(255, 255, 255))

    run = p.add_run()
    try:
        run.add_picture(asset_path, width=Inches(target_width_inches))
    except Exception as e:
        print(f"[Warning] Failed to insert image picture: {e}")

    # In high-fidelity mode, subtle label caption (omit for portraits, profile photos, and avatars)
    img_lbl = str(block.get("image_label", "")).lower()
    if mode == "high_fidelity" and img_lbl and not any(k in img_lbl for k in ("photo", "profile", "avatar", "portrait")):
        p_cap = container.add_paragraph()
        p_cap.alignment = p.alignment
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(4)
        run_cap = p_cap.add_run(f"[{block.get('image_label').upper()}]")
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(7.5)
        run_cap.font.color.rgb = RGBColor(148, 163, 184)

DATE_RANGE_REGEX = re.compile(
    r'(?:'
    r'(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+)?(?:19|20)\d{2}\s*[-–—/]\s*'
    r'(?:(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+)?(?:19|20)\d{2}|PRESENT|CURRENT|ONGOING|TODAY)'
    r'|'
    r'(?:19|20)\d{2}\s*[-–—/]\s*(?:19|20)\d{2}'
    r')',
    re.IGNORECASE
)

def render_experience_header(container, lines: list[str], date_str: str, available_width_inches: float = 5.0, accent_color=None, default_text_color=(15, 23, 42)):
    """
    Renders a job title/employer on the left and dates right-aligned across from them
    using a 1-row, 2-cell borderless table with fixed layout.
    """
    tbl = container.add_table(rows=1, cols=2)
    remove_table_borders(tbl)
    tbl._tbl.tblPr.append(parse_xml(r'<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))

    date_w = min(1.8, max(1.2, available_width_inches * 0.30))
    left_w = max(2.0, available_width_inches - date_w)

    cell_left = tbl.cell(0, 0)
    cell_right = tbl.cell(0, 1)
    cell_left.width = Inches(left_w)
    cell_right.width = Inches(date_w)

    set_cell_margins(cell_left, top=30, bottom=10, left=0, right=30)
    set_cell_margins(cell_right, top=30, bottom=10, left=30, right=0)

    p_l = cell_left.paragraphs[0]
    p_l.paragraph_format.line_spacing = 1.05
    p_l.paragraph_format.space_before = Pt(4)
    p_l.paragraph_format.space_after = Pt(1)

    if lines:
        r0 = p_l.add_run(lines[0])
        r0.font.name = "Calibri"
        r0.font.size = Pt(10.5)
        r0.bold = True
        r0.font.color.rgb = RGBColor(*default_text_color)

    for l in lines[1:]:
        p_sub = cell_left.add_paragraph()
        p_sub.paragraph_format.line_spacing = 1.05
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1)
        r_sub = p_sub.add_run(l)
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(9.5)
        r_sub.italic = True
        r_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_r = cell_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.line_spacing = 1.05
    p_r.paragraph_format.space_before = Pt(4)
    p_r.paragraph_format.space_after = Pt(1)

    r_d = p_r.add_run(date_str)
    r_d.font.name = "Calibri"
    r_d.font.size = Pt(9.5)
    r_d.bold = True
    if accent_color:
        r_d.font.color.rgb = hex_to_rgb(accent_color, default=default_text_color)
    else:
        r_d.font.color.rgb = RGBColor(71, 85, 105)

def render_reference_card(container, element: dict, available_width_inches: float = 2.5, accent_color=None):
    """
    Renders an individual reference entry inside a 2-column reference table cell.
    Formats the name in bold and subsequent title/contact info in compact typography.
    """
    raw_text = str(element.get("text") or element.get("content") or "").strip()
    if not raw_text:
        return
    lines = [l.strip() for l in raw_text.split('\n') if l.strip()]
    if not lines:
        return

    p = container.add_paragraph()
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)

    r_name = p.add_run(lines[0])
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(10)
    r_name.bold = True
    r_name.font.color.rgb = RGBColor(30, 41, 59)

    for line in lines[1:]:
        p_sub = container.add_paragraph()
        p_sub.paragraph_format.line_spacing = 1.05
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1)
        r_sub = p_sub.add_run(line)
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(9.0)
        if any(c in line for c in ["CTO", "CEO", "Manager", "Director", "President", "Lead", "/"]):
            r_sub.italic = True
            r_sub.font.color.rgb = RGBColor(71, 85, 105)
        else:
            r_sub.font.color.rgb = RGBColor(51, 65, 85)

def render_table_block(container, block, available_width_inches=7.5):
    """
    Renders an extracted table into a native editable Word table with cell borders,
    background fills, and alignments.
    """
    table_rows = block.get("table_rows", [])
    if not table_rows:
        return

    num_rows = len(table_rows)
    num_cols = max(len(row) for row in table_rows) if num_rows > 0 else 0
    if num_rows == 0 or num_cols == 0:
        return

    table = container.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_bg = clean_hex_color(block.get("table_header_bg"), default="F1F5F9")
    has_borders = block.get("table_has_borders", True)

    if has_borders:
        set_table_borders(table, border_color="CBD5E1", inside_h_color="E2E8F0")
    else:
        remove_table_borders(table)

    col_width = Inches(available_width_inches / num_cols)

    for r_idx, row_data in enumerate(table_rows):
        row = table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))

        is_header = (r_idx == 0)
        if is_header:
            trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell.width = col_width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

            cell_text = str(row_data[c_idx]).strip() if c_idx < len(row_data) else ""

            tcPr = cell._tc.get_or_add_tcPr()
            if is_header and header_bg:
                shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), header_bg))
                tcPr.append(shd)
            elif r_idx % 2 == 1:
                shd = parse_xml(r'<w:shd %s w:fill="F8FAFC"/>' % nsdecls('w'))
                tcPr.append(shd)

            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            # Detect right-aligned numeric cells
            if cell_text.replace('$', '').replace('%', '').replace(',', '').replace('.', '').strip().isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(cell_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(10 if is_header else 9.5)
            run.bold = is_header
            if is_header:
                run.font.color.rgb = RGBColor(30, 41, 59)
            else:
                run.font.color.rgb = RGBColor(51, 65, 85)

    p_spacer = container.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(3)
    p_spacer.paragraph_format.space_after = Pt(5)

def render_text_block(container, block, default_text_color=(15, 23, 42), accent_color=None):
    """
    Renders heading, paragraph, or bullet block preserving typography, alignment,
    colors, background cards, and contact icons.
    """
    # Accept either 'text' or 'content' key
    content = str(block.get("text") or block.get("content") or "").strip()
    if not content:
        return

    is_heading = block.get("is_heading", False) or block.get("type") == "heading"
    is_bullet = block.get("type") == "bullet" or block.get("is_bullet", False)
    icon = block.get("icon_type", "")

    font_size = float(block.get("font_size_pt", 13.0 if is_heading else 10.0))
    is_bold = block.get("is_bold", is_heading)
    is_italic = block.get("is_italic", False)

    detected_color = block.get("text_color")
    text_color = hex_to_rgb(detected_color, default=default_text_color) if detected_color else RGBColor(*default_text_color)

    p = container.add_paragraph()
    p.alignment = get_alignment(block.get("text_align", "left"))
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.space_before = Pt(5 if is_heading else 1.5)
    p.paragraph_format.space_after = Pt(2.5 if is_heading else 2)

    # Shaded paragraph background if element has card bg_color
    bg_color = block.get("bg_color")
    if bg_color and isinstance(bg_color, str) and bg_color.strip():
        clean_bg = clean_hex_color(bg_color)
        pPr = p._p.get_or_add_pPr()
        shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), clean_bg))
        pPr.append(shd)

    # Accent dot before headings if accent_color is present
    if is_heading and not is_bullet and accent_color:
        r_dot = p.add_run("● ")
        r_dot.font.name = "Calibri"
        r_dot.font.size = Pt(font_size)
        r_dot.font.color.rgb = hex_to_rgb(accent_color, default=(229, 169, 20))

    # Icon prefix for contact info
    if icon:
        icon_symbol = ""
        if icon == "phone":
            icon_symbol = "☎  "
        elif icon in ("email", "website"):
            icon_symbol = "✉  " if "@" in content else "🌐  "
        elif icon == "location":
            icon_symbol = "⚲  "
        elif icon in ("bullet_dot", "star"):
            icon_symbol = "●  "
        if icon_symbol:
            r_ic = p.add_run(icon_symbol)
            r_ic.font.name = "Segoe UI Symbol"
            r_ic.font.size = Pt(font_size)
            r_ic.font.color.rgb = hex_to_rgb(accent_color, default=default_text_color) if accent_color else text_color

    if is_bullet:
        p.style = 'List Bullet'
        cleaned_content = content.lstrip('•-*» ').strip()
        run = p.add_run(cleaned_content)
    else:
        run = p.add_run(content)

    run.font.name = 'Calibri'
    run.font.size = Pt(font_size)
    run.bold = is_bold
    run.italic = is_italic
    run.font.color.rgb = text_color

    # Decorative underline bar beneath section headings
    if is_heading and not is_bullet:
        is_light_text = (default_text_color[0] > 200)
        underline_hex = "FFFFFF" if is_light_text else clean_hex_color(accent_color or detected_color, default="17304B")
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            r'<w:pBdr %s>'
            r'  <w:bottom w:val="single" w:sz="8" w:space="2" w:color="%s"/>'
            r'</w:pBdr>' % (nsdecls('w'), underline_hex)
        )
        pPr.append(pBdr)
        p.paragraph_format.space_after = Pt(3)

def render_element(container, item: dict, available_width_inches: float = 7.5, mode: str = "editable", default_text_color=(15, 23, 42), max_image_width=2.5, accent_color=None):
    """
    Dispatches rendering based on element type.
    """
    item_type = item.get("type", "paragraph")

    if item_type == "image":
        render_image_block(container, item, available_width_inches=available_width_inches, mode=mode, max_image_width=max_image_width)
    elif item_type == "table" or item.get("table_rows"):
        render_table_block(container, item, available_width_inches=available_width_inches)
    elif item_type == "skill_bar":
        render_skill_bars_grid(container, [item], available_width_inches=available_width_inches, accent_hex=accent_color or "E5A914")
    elif item_type == "divider":
        render_divider(container, item, available_width_inches=available_width_inches)
    else:
        render_text_block(container, item, default_text_color=default_text_color, accent_color=accent_color)

def render_single_page(doc, page_data, mode: str = "editable"):
    """
    Renders one page preserving its exact detected layout structure:
    - If page has a sidebar (CVs, profile resumes): renders two-column layout grid with colored sidebar
    - If page does not have a sidebar (invoices, forms, reports): renders standard document flow
    """
    if isinstance(page_data, list):
        elements = page_data
        has_sidebar = any(e.get("column") in ["left", "right"] for e in elements)
        sidebar_pos = "left" if any(e.get("column") == "left" for e in elements) else "none"
        sidebar_bg = None
        sidebar_pct = 33
    elif isinstance(page_data, dict):
        elements = page_data.get("elements", [])
        has_sidebar = page_data.get("has_sidebar", False) or any(e.get("column") in ["left", "right"] for e in elements)
        sidebar_pos = page_data.get("sidebar_position", "left")
        sidebar_bg = page_data.get("sidebar_bg_color")
        sidebar_pct = page_data.get("sidebar_width_percent", 34)
    else:
        return

    if not elements:
        return

    # If column assignments are missing, infer them from bounding boxes if present
    if not any(e.get("column") in ["left", "right"] for e in elements):
        left_cand = [e for e in elements if len(e.get("bbox", [])) == 4 and e["bbox"][3] <= 420 and e["bbox"][1] < 120]
        right_cand = [e for e in elements if len(e.get("bbox", [])) == 4 and e["bbox"][1] >= 340]
        has_full_width_body = any(
            len(e.get("bbox", [])) == 4 and e["bbox"][1] < 150 and e["bbox"][3] > 800 and e["bbox"][0] > 180
            for e in elements
        )
        if len(left_cand) >= 2 and right_cand and not has_full_width_body:
            max_left_y = max(e["bbox"][2] for e in left_cand)
            if max_left_y >= 350:
                has_sidebar = True
                sidebar_pos = "left"
                left_maxs = [e["bbox"][3] for e in left_cand]
                if left_maxs:
                    avg_w = int(sum(left_maxs) / len(left_maxs) / 10)
                    sidebar_pct = max(25, min(45, avg_w))
                for e in elements:
                    bbox = e.get("bbox", [])
                    if len(bbox) == 4:
                        ymin, xmin, ymax, xmax = bbox
                        if xmin < 150 and xmax > 850:
                            e["column"] = "single"
                        elif xmax <= 420:
                            e["column"] = "left"
                        elif xmin >= 340:
                            e["column"] = "right"
                        else:
                            e["column"] = "right"

    left_elements = [e for e in elements if e.get("column") == "left"]
    right_elements = [e for e in elements if e.get("column") == "right"]
    single_elements = [e for e in elements if e.get("column") not in ["left", "right"]]

    # If document has sidebar or multi-column layout
    if has_sidebar and (left_elements or right_elements or elements):
        SIDEBAR_SEC_KEYS = [
            "CONTACT", "CONTACT ME", "EDUCATION", "SKILLS", "SKILL", "LANGUAGES", "LANGUAGE",
            "EXPERTISE", "HOBBIES", "INTERESTS", "AWARDS", "CERTIFICATIONS"
        ]
        MAIN_SEC_KEYS = [
            "PROFILE", "ABOUT", "ABOUT ME", "SUMMARY", "WORK EXPERIENCE", "EXPERIENCE",
            "EMPLOYMENT", "REFERENCE", "REFERENCES", "PROJECTS", "PUBLICATIONS"
        ]

        new_left = []
        new_right = []
        top_single = []

        # Find profile photo if present to place at top of sidebar
        photo_el = None
        for e in elements:
            if e.get("type") == "image" and any(l in str(e.get("image_label", "")).lower() for l in ["photo", "profile", "avatar", "portrait"]):
                photo_el = e
                break

        current_col = "right"  # Elements before any section heading (Name, Title) belong to right
        for e in elements:
            if e is photo_el:
                continue

            bbox = e.get("bbox", [])
            # Top full-width banner
            if len(bbox) == 4 and bbox[2] < 160 and bbox[1] < 150 and bbox[3] > 800:
                top_single.append(e)
                continue

            # 1. Respect explicit column assignment from OCR/Gemini
            col = e.get("column")
            if col in ["left", "right"]:
                target_col = col
                current_col = col
            else:
                # 2. Fallback if column assignment is missing: check bbox or heading keywords
                if len(bbox) == 4:
                    if bbox[3] <= (sidebar_pct * 10 + 30):
                        target_col = "left"
                    else:
                        target_col = "right"
                    current_col = target_col
                else:
                    txt = str(e.get("text") or e.get("content") or "").strip().upper()
                    is_hdr = (e.get("type") == "heading" or e.get("is_heading"))
                    if is_hdr:
                        if any(k in txt for k in SIDEBAR_SEC_KEYS):
                            current_col = "left"
                        elif any(k in txt for k in MAIN_SEC_KEYS):
                            current_col = "right"
                        elif e.get("font_size_pt", 0) >= 18:
                            current_col = "right"
                    target_col = current_col

            if target_col == "left":
                new_left.append(e)
            else:
                new_right.append(e)

        if photo_el:
            new_left.insert(0, photo_el)

        left_elements = filter_redundant_dividers(new_left)
        right_elements = filter_redundant_dividers(new_right)

        # Determine background color for sidebar
        if not sidebar_bg:
            bg_candidates = [e.get("bg_color") for e in left_elements if e.get("bg_color")]
            sidebar_bg = bg_candidates[0] if bg_candidates else "1C3144"
        clean_bg = clean_hex_color(sidebar_bg, default="1C3144")

        for item in top_single:
            render_element(doc, item, available_width_inches=7.8, mode=mode)

        # 2. Render 2-Column Sidebar Layout Grid
        grid = doc.add_table(rows=1, cols=2)
        remove_table_borders(grid)
        grid._tbl.tblPr.append(parse_xml(r'<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))

        trPr = grid.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        trPr.append(parse_xml(r'<w:trHeight %s w:val="14200" w:hRule="atLeast"/>' % nsdecls('w')))

        col_ratio = max(0.25, min(0.48, float(sidebar_pct) / 100.0))
        left_w = 7.8 * col_ratio
        right_w = 7.8 - left_w

        left_cell = grid.cell(0, 0)
        right_cell = grid.cell(0, 1)

        left_cell.width = Inches(left_w)
        right_cell.width = Inches(right_w)

        # Apply detected background shading to sidebar cell
        shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), clean_bg))
        left_cell._tc.get_or_add_tcPr().append(shd)

        # Set cell padding: compact to fit 1 page
        set_cell_margins(left_cell, top=80, bottom=80, left=100, right=100)
        set_cell_margins(right_cell, top=80, bottom=80, left=140, right=100)

        accent_color = page_data.get("accent_color") if isinstance(page_data, dict) else None
        header_bg_color = page_data.get("header_bg_color") if isinstance(page_data, dict) else None
        has_corners = page_data.get("has_corner_decorations", False) if isinstance(page_data, dict) else False

        # Render Left Sidebar Elements
        first_left = True

        # 1. Top Corner Decoration if present
        if has_corners and accent_color:
            corner_asset = generate_corner_triangle_asset(accent_color, corner="top_left", size_px=120)
            if corner_asset and os.path.exists(corner_asset):
                p_c = left_cell.paragraphs[0]
                p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_c.paragraph_format.space_before = Pt(0)
                p_c.paragraph_format.space_after = Pt(2)
                r_c = p_c.add_run()
                r_c.add_picture(corner_asset, width=Inches(0.75))
                first_left = False

        for item in left_elements:
            # Clear initial empty paragraph on first write
            if first_left and len(left_cell.paragraphs) > 0 and not left_cell.paragraphs[0].text:
                p = left_cell.paragraphs[0]
                first_left = False
                # If first item is an image, place it inside the cell
                if item.get("type") == "image" and item.get("image_asset_path") and os.path.exists(item["image_asset_path"]):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(6)
                    # Apply circular mask for profile photos in sidebar
                    asset_path = item["image_asset_path"]
                    label = str(item.get("image_label", "")).lower()
                    if "photo" in label or "profile" in label:
                        asset_path = make_circular_image(asset_path, border_width=6, border_color=(255, 255, 255))
                    run = p.add_run()
                    run.add_picture(asset_path, width=Inches(min(left_w - 0.4, 1.35)))
                    continue
                else:
                    p._p.getparent().remove(p._p)

            render_element(
                left_cell, item,
                available_width_inches=left_w - 0.3,
                mode=mode,
                default_text_color=(255, 255, 255),  # Default white text on dark sidebar
                max_image_width=min(left_w - 0.4, 1.35),
                accent_color=accent_color
            )

        # Render Right Main Body Elements
        first_right = True
        # Detect REFERENCE heading index to handle two-column references in the main body
        ref_heading_idx = None
        for ri, item in enumerate(right_elements):
            if item.get("type") == "heading" and "REFERENCE" in str(item.get("text") or item.get("content") or "").upper():
                # Verify this REFERENCE section is a trailing section, not followed by other major body sections
                subsequent_headings = [
                    str(s.get("text") or s.get("content") or "").upper()
                    for s in right_elements[ri + 1:]
                    if (s.get("type") == "heading" or s.get("is_heading"))
                ]
                has_major_section_after = any(
                    any(k in sh for k in ["EXPERIENCE", "WORK", "ABOUT", "PROFILE", "EDUCATION", "SKILL"])
                    for sh in subsequent_headings
                )
                if not has_major_section_after:
                    ref_heading_idx = ri
                break

        # Elements before the REFERENCE heading (or all if no reference)
        main_right = right_elements[:ref_heading_idx] if ref_heading_idx is not None else right_elements
        ref_elements = right_elements[ref_heading_idx:] if ref_heading_idx is not None else []

        # Check for header banner card (candidate name and title)
        name_el = None
        title_el = None
        start_idx = 0
        if len(main_right) > 0 and (main_right[0].get("type") == "heading" or main_right[0].get("font_size_pt", 0) >= 18):
            first_txt = str(main_right[0].get("text") or main_right[0].get("content") or "").upper()
            if header_bg_color or (main_right[0].get("font_size_pt", 0) >= 18 and "CV" not in first_txt):
                name_el = main_right[0]
                start_idx = 1
                if len(main_right) > 1 and main_right[1].get("type") in ("paragraph", "heading"):
                    second_txt = str(main_right[1].get("text") or main_right[1].get("content") or "").upper()
                    if not any(s in second_txt for s in ["PROFILE", "ABOUT", "EXPERIENCE", "WORK", "EDUCATION", "SKILL", "CONTACT", "REFERENCE", "SUMMARY", "OBJECTIVE"]):
                        title_el = main_right[1]
                        start_idx = 2

        if name_el:
            if first_right and len(right_cell.paragraphs) > 0 and not right_cell.paragraphs[0].text:
                first_right = False
                p = right_cell.paragraphs[0]
                p._p.getparent().remove(p._p)

            h_tbl = right_cell.add_table(rows=1, cols=1)
            remove_table_borders(h_tbl)
            h_tbl._tbl.tblPr.append(parse_xml(r'<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
            h_cell = h_tbl.cell(0, 0)
            h_cell.width = Inches(right_w - 0.3)
            set_cell_margins(h_cell, top=80, bottom=80, left=120, right=100)

            # Determine header background shading: default to light-grey (#F1F3F5) if white or omitted
            hdr_hex = clean_hex_color(header_bg_color, default="F1F3F5")
            if hdr_hex in ("FFFFFF", "FFF"):
                hdr_hex = "F1F3F5"
            shd_h = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), hdr_hex))
            h_cell._tc.get_or_add_tcPr().append(shd_h)

            # Left accent border: default to gold (#E5A914) if white or omitted
            acc_hex = clean_hex_color(accent_color, default="E5A914") if accent_color else "E5A914"
            if acc_hex in ("FFFFFF", "FFF"):
                acc_hex = "E5A914"
            tcBorders = parse_xml(r'<w:tcBorders %s><w:left w:val="single" w:sz="24" w:space="0" w:color="%s"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>' % (nsdecls('w'), acc_hex))
            h_cell._tc.get_or_add_tcPr().append(tcBorders)

            p_name = h_cell.paragraphs[0]
            p_name.paragraph_format.line_spacing = 1.0
            p_name.paragraph_format.space_before = Pt(1)
            p_name.paragraph_format.space_after = Pt(1)
            name_text = str(name_el.get("text") or name_el.get("content") or "").strip()
            r_n = p_name.add_run(name_text)
            r_n.font.name = "Calibri"
            r_n.font.size = Pt(20)
            r_n.bold = True
            r_n.font.color.rgb = RGBColor(30, 41, 59)

            if title_el:
                title_text = str(title_el.get("text") or title_el.get("content") or "").strip()
                p_t = h_cell.add_paragraph()
                p_t.paragraph_format.line_spacing = 1.0
                p_t.paragraph_format.space_before = Pt(1)
                p_t.paragraph_format.space_after = Pt(1)
                r_t = p_t.add_run(title_text)
                r_t.font.name = "Calibri"
                r_t.font.size = Pt(10.5)
                r_t.font.color.rgb = RGBColor(100, 116, 139)

            main_right = main_right[start_idx:]

        # Separate skill items to render in a grid
        skill_items = [e for e in main_right if e.get("type") == "skill_bar"]
        body_items = [e for e in main_right if e.get("type") != "skill_bar"]

        idx_body = 0
        while idx_body < len(body_items):
            item = body_items[idx_body]
            item_text = str(item.get("text") or item.get("content") or "").strip()
            item_type = item.get("type", "paragraph")
            is_sec_heading = (item_type == "heading" or item.get("is_heading")) and any(
                s in item_text.upper() for s in ["EXPERIENCE", "PROFILE", "EDUCATION", "REFERENCE", "SUMMARY", "ABOUT", "SKILL", "CONTACT"]
            )

            # Check if this item is a job title/experience header with a date range
            date_match = DATE_RANGE_REGEX.search(item_text) if item_type in ("paragraph", "heading") else None
            next_is_date = (
                not date_match
                and not is_sec_heading
                and idx_body + 1 < len(body_items)
                and DATE_RANGE_REGEX.fullmatch(str(body_items[idx_body + 1].get("text") or body_items[idx_body + 1].get("content") or "").strip())
            )

            if date_match and not is_sec_heading and not item_text.startswith("●"):
                date_str = date_match.group(0).strip()
                clean_text = item_text[:date_match.start()] + item_text[date_match.end():]
                clean_text = re.sub(r'[\(\)]', '', clean_text)
                clean_text = re.sub(r'\s*\|\s*', '\n', clean_text)
                lines = [l.strip(' \t,-|') for l in clean_text.split('\n') if l.strip(' \t,-|')]

                if first_right and len(right_cell.paragraphs) > 0 and not right_cell.paragraphs[0].text:
                    first_right = False
                    p = right_cell.paragraphs[0]
                    p._p.getparent().remove(p._p)

                render_experience_header(
                    right_cell, lines, date_str,
                    available_width_inches=right_w - 0.3,
                    accent_color=accent_color,
                    default_text_color=(15, 23, 42)
                )
                idx_body += 1
                continue
            elif next_is_date and not item_text.startswith("●"):
                date_str = str(body_items[idx_body + 1].get("text") or body_items[idx_body + 1].get("content") or "").strip()
                lines = [l.strip(' \t,-|') for l in item_text.split('\n') if l.strip(' \t,-|')]

                if first_right and len(right_cell.paragraphs) > 0 and not right_cell.paragraphs[0].text:
                    first_right = False
                    p = right_cell.paragraphs[0]
                    p._p.getparent().remove(p._p)

                render_experience_header(
                    right_cell, lines, date_str,
                    available_width_inches=right_w - 0.3,
                    accent_color=accent_color,
                    default_text_color=(15, 23, 42)
                )
                idx_body += 2
                continue

            if first_right and len(right_cell.paragraphs) > 0 and not right_cell.paragraphs[0].text:
                first_right = False
                p = right_cell.paragraphs[0]
                p._p.getparent().remove(p._p)

            render_element(
                right_cell, item,
                available_width_inches=right_w - 0.3,
                mode=mode,
                default_text_color=(15, 23, 42),
                max_image_width=right_w - 0.5,
                accent_color=accent_color
            )
            idx_body += 1

        if skill_items:
            render_skill_bars_grid(right_cell, skill_items, available_width_inches=right_w - 0.4, accent_hex=accent_color or "E5A914")

        # Bottom-right corner shape if present
        if has_corners and accent_color:
            corner_br = generate_corner_triangle_asset(accent_color, corner="bottom_right", size_px=120)
            if corner_br and os.path.exists(corner_br):
                p_br = right_cell.add_paragraph()
                p_br.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p_br.paragraph_format.space_before = Pt(4)
                p_br.paragraph_format.space_after = Pt(0)
                r_br = p_br.add_run()
                r_br.add_picture(corner_br, width=Inches(0.75))

        # Render REFERENCE section with two-column sub-layout if applicable
        if ref_elements:
            # Render the REFERENCE heading itself
            ref_heading = ref_elements[0]
            if first_right and len(right_cell.paragraphs) > 0 and not right_cell.paragraphs[0].text:
                first_right = False
                p = right_cell.paragraphs[0]
                p._p.getparent().remove(p._p)
            render_element(
                right_cell, ref_heading,
                available_width_inches=right_w - 0.3,
                mode=mode,
                default_text_color=(15, 23, 42),
                max_image_width=right_w - 0.5,
                accent_color=accent_color
            )

            # Split remaining reference entries into left-half and right-half
            ref_body = ref_elements[1:]
            if ref_body:
                # 1. Determine column midpoint from bbox data if available
                all_xmins = [e.get("bbox", [0, 0, 0, 0])[1] for e in ref_body if len(e.get("bbox", [])) >= 4]
                ref_left = []
                ref_right = []
                ref_no_bbox = []
                if all_xmins and (max(all_xmins) - min(all_xmins)) > 100:
                    midpoint = (min(all_xmins) + max(all_xmins)) / 2.0
                    ref_left = [e for e in ref_body if len(e.get("bbox", [])) >= 4 and e["bbox"][1] < midpoint]
                    ref_right = [e for e in ref_body if len(e.get("bbox", [])) >= 4 and e["bbox"][1] >= midpoint]
                    ref_no_bbox = [e for e in ref_body if len(e.get("bbox", [])) < 4]

                # 2. Fallback: If bounding boxes lack horizontal spread or are missing,
                # partition ref_body elements into two columns (left: Estelle Darcy, right: Harper Richard)
                if not (ref_left and ref_right):
                    if len(ref_body) >= 2:
                        mid = (len(ref_body) + 1) // 2
                        ref_left = ref_body[:mid]
                        ref_right = ref_body[mid:]
                        ref_no_bbox = []
                    elif len(ref_body) == 1:
                        single_txt = str(ref_body[0].get("text") or ref_body[0].get("content") or "").strip()
                        parts = [p.strip() for p in re.split(r'\n\s*\n', single_txt) if p.strip()]
                        if len(parts) >= 2:
                            el_base = dict(ref_body[0])
                            ref_left = [dict(el_base, text=parts[0], content=parts[0])]
                            ref_right = [dict(el_base, text=parts[1], content=parts[1])]
                            ref_no_bbox = []

                if ref_left and ref_right:
                    # Create a 2-column sub-table for references
                    ref_table = right_cell.add_table(rows=1, cols=2)
                    remove_table_borders(ref_table)
                    ref_table._tbl.tblPr.append(parse_xml(r'<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))

                    half_w = (right_w - 0.3) / 2.0
                    ref_table.cell(0, 0).width = Inches(half_w)
                    ref_table.cell(0, 1).width = Inches(half_w)

                    ref_left_cell = ref_table.cell(0, 0)
                    ref_right_cell = ref_table.cell(0, 1)

                    set_cell_margins(ref_left_cell, top=20, bottom=20, left=0, right=40)
                    set_cell_margins(ref_right_cell, top=20, bottom=20, left=40, right=0)

                    # Clear initial empty paragraphs
                    for c in (ref_left_cell, ref_right_cell):
                        if c.paragraphs and not c.paragraphs[0].text:
                            c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)

                    for e in ref_left:
                        render_reference_card(ref_left_cell, e, available_width_inches=half_w - 0.1, accent_color=accent_color)
                    for e in ref_right:
                        render_reference_card(ref_right_cell, e, available_width_inches=half_w - 0.1, accent_color=accent_color)
                    for e in ref_no_bbox:
                        render_element(right_cell, e, available_width_inches=right_w - 0.3,
                                       mode=mode, default_text_color=(15, 23, 42), max_image_width=right_w - 0.5,
                                       accent_color=accent_color)
                else:
                    for e in ref_body:
                        render_element(right_cell, e, available_width_inches=right_w - 0.3,
                                       mode=mode, default_text_color=(15, 23, 42), max_image_width=right_w - 0.5,
                                       accent_color=accent_color)

        # Trailing paragraph after layout table: format with 1pt to avoid pushing blank page 2 in Word
        p_tail = doc.add_paragraph()
        p_tail.paragraph_format.space_before = Pt(0)
        p_tail.paragraph_format.space_after = Pt(0)
        p_tail.paragraph_format.line_spacing = Pt(1)
        r_tail = p_tail.add_run()
        r_tail.font.size = Pt(1)

    else:
        # Standard Single-Column Document Flow (Invoices, Letters, Reports, Forms)
        for item in elements:
            render_element(doc, item, available_width_inches=7.8, mode=mode)

def build_advanced_docx(pages_ir, output_path: str, mode: str = "editable"):
    """
    Constructs a Microsoft Word document (.docx) with visual document reconstruction:
    - Preserves sidebars with detected colors and widths
    - Embeds photos, logos, and signatures in their correct relative positions
    - Preserves native editable tables, headers, and cell formatting
    - Handles multi-page documents with native page breaks
    """
    doc = Document()

    # Apply 0.35-inch margins and standard Letter size (8.5 x 11.0 in)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.35)
        section.right_margin = Inches(0.35)

    if not pages_ir:
        doc.add_paragraph("No content was extracted from the document.")
        doc.save(output_path)
        return

    # Normalize pages_ir to list of pages
    if isinstance(pages_ir, dict):
        pages_list = [pages_ir]
    elif isinstance(pages_ir, list) and len(pages_ir) > 0 and isinstance(pages_ir[0], dict) and "elements" not in pages_ir[0]:
        # Flat list of elements
        pages_list = [pages_ir]
    else:
        pages_list = pages_ir

    total_pages = len(pages_list)

    for page_idx, page_data in enumerate(pages_list):
        render_single_page(doc, page_data, mode=mode)

        # Insert native page break between consecutive pages
        if page_idx < total_pages - 1:
            doc.add_page_break()

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)